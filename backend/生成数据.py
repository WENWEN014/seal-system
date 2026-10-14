# -*- coding: utf-8 -*-
"""
反透视：把宽表 Excel 转成 3 个 CSV 长表
- items.csv      : 最小数据单元，一行一枚章/一份执照（1542 行）
- locations.csv  : 位置与人员的统一节点表（支持归属关系）
- movements.csv  : 流转轨迹（每次移动留痕）

输出目录：环境变量 SEAL_DATA_DIR，默认 E:\证照印章数据
源表：Z:\ceshi\刻章情况.xlsx（只读）
"""
import os, sys, csv, json
from datetime import datetime, date
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SRC = Path(r"Z:\ceshi\刻章情况.xlsx")
OUT_DIR = Path(os.environ.get("SEAL_DATA_DIR", r"E:\证照印章数据"))

SEALS = ["公章", "法人章", "财务章", "合同章", "发票章"]
STAMP = "2026-09-30 11:40"

def dval(v):
    if v is None: return ""
    if isinstance(v, (datetime, date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()

def main():
    if not SRC.exists():
        print(f"[FATAL] 找不到 {SRC}")
        return 1

    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["刻章情况"]

    # ---------- 1. 读出库明细：哪些章当前在向向处 ----------
    out_set = set()  # (公司名, 章类型)
    out_rows_data = []
    if "出库明细" in wb.sheetnames:
        ws_out = wb["出库明细"]
        for r in range(3, ws_out.max_row + 1):
            name = dval(ws_out.cell(r, 5).value)   # 公司名称在第5列
            if not name: continue
            seals_here = []
            for i, s in enumerate(SEALS):
                if dval(ws_out.cell(r, 8 + i).value) in ("1", "1.0"):  # 章在第8-12列
                    out_set.add((name, s))
                    seals_here.append(s)
            out_rows_data.append({"name": name, "seals": seals_here,
                                  "family": dval(ws_out.cell(r, 4).value),   # 家族第4列
                                  "legal": dval(ws_out.cell(r, 6).value)})   # 法人第6列

    # ---------- 2. 反透视：生成 items ----------
    items = []
    seq = 0
    for r in range(3, ws.max_row + 1):
        name = dval(ws.cell(r, 4).value)
        if not name: continue
        family = dval(ws.cell(r, 3).value)
        legal = dval(ws.cell(r, 5).value)
        ctype = dval(ws.cell(r, 1).value) or "大陆"
        reg = dval(ws.cell(r, 2).value)
        pos = dval(ws.cell(r, 12).value)

        # 章
        for i, s in enumerate(SEALS):
            if dval(ws.cell(r, 7 + i).value) in ("1", "1.0"):
                seq += 1
                is_out = (name, s) in out_set
                items.append({
                    "item_id": f"IT{seq:05d}",
                    "公司名称": name,
                    "家族": family,
                    "法人": legal,
                    "公司类型": ctype,
                    "注册时间": reg,
                    "资料类型": s,
                    "资料类别": "章",
                    "当前位置ID": "P001" if is_out else "",
                    "当前位置": "向向" if is_out else pos,
                    "当前位置类型": "人员" if is_out else "位置",
                    "状态": "在外部" if is_out else "在库",
                    "持有人": "向向" if is_out else "",
                    "出库时间": "" if is_out else "",
                    "更新时间": STAMP,
                    "备注": "存量补录（原出库明细表）" if is_out else "",
                })
        # 营业执照：每公司一份，位置待分配（柜二）
        seq += 1
        items.append({
            "item_id": f"IT{seq:05d}",
            "公司名称": name,
            "家族": family,
            "法人": legal,
            "公司类型": ctype,
            "注册时间": reg,
            "资料类型": "营业执照",
            "资料类别": "执照",
            "当前位置ID": "",
            "当前位置": "待分配",
            "当前位置类型": "位置",
            "状态": "在库",
            "持有人": "",
            "出库时间": "",
            "更新时间": STAMP,
            "备注": "存放柜二，位置待补录",
        })

    # ---------- 3. 位置与人员统一节点表（含 4 排 × 3 位坐标）----------
    ROWS_PER_CAB, COLS_PER_ROW = 4, 3
    locations = []
    lid = 0
    cab_counter = {}   # 每柜已放筐数，用于计算 排/位
    # 柜子分布表 → 位置节点
    if "柜子分布" in wb.sheetnames:
        ws_cab = wb["柜子分布"]
        cur_cab = None
        for r in range(1, ws_cab.max_row + 1):
            cab = dval(ws_cab.cell(r, 2).value)
            bsk = dval(ws_cab.cell(r, 3).value)
            fam = dval(ws_cab.cell(r, 4).value)
            if cab: cur_cab = cab
            if bsk:
                lid += 1
                n = cab_counter.get(cur_cab, 0)
                row_no = n // COLS_PER_ROW + 1
                col_no = n % COLS_PER_ROW + 1
                cab_counter[cur_cab or ""] = n + 1
                locations.append({
                    "id": f"L{lid:03d}", "名称": bsk, "类型": "位置",
                    "柜子": cur_cab or "", "排": str(row_no), "位": str(col_no),
                    "关联家族": fam, "负责人": "", "备注": "",
                })
    # 人员节点
    locations.append({"id": "P001", "名称": "向向", "类型": "人员",
                      "柜子": "", "排": "", "位": "",
                      "关联家族": "", "负责人": "", "备注": "外部办理人"})

    # ---------- 4. 初始流转轨迹 ----------
    movements = []
    mid = 0
    # 存量出库：位置 → 向向
    for row in out_rows_data:
        for s in row["seals"]:
            it = next((x for x in items if x["公司名称"] == row["name"]
                       and x["资料类型"] == s and x["状态"] == "在外部"), None)
            if not it: continue
            mid += 1
            movements.append({
                "事件ID": f"MV{mid:05d}", "时间": STAMP, "item_id": it["item_id"],
                "公司名称": row["name"], "资料类型": s, "动作": "出库",
                "起点ID": "", "起点名称": "（原位置未记录）", "起点类型": "位置",
                "终点ID": "P001", "终点名称": "向向", "终点类型": "人员",
                "操作人": "系统初始化", "事由": "存量补录（原出库明细表，日期未记录）", "备注": "",
            })

    # ---------- 5. 写 CSV ----------
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    def write_csv(fn, rows, headers):
        p = OUT_DIR / fn
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=headers)
            w.writeheader()
            for r in rows: w.writerow(r)
        print(f"  {fn}: {len(rows)} 行 → {p}")

    print(f"输出目录：{OUT_DIR}")
    write_csv("items.csv", items,
              ["item_id", "公司名称", "家族", "法人", "公司类型", "注册时间",
               "资料类型", "资料类别", "当前位置ID", "当前位置", "当前位置类型",
               "状态", "持有人", "出库时间", "更新时间", "备注"])
    write_csv("locations.csv", locations,
              ["id", "名称", "类型", "柜子", "排", "位", "关联家族", "负责人", "备注"])
    write_csv("movements.csv", movements,
              ["事件ID", "时间", "item_id", "公司名称", "资料类型", "动作",
               "起点ID", "起点名称", "起点类型", "终点ID", "终点名称", "终点类型",
               "操作人", "事由", "备注"])

    # 统计
    from collections import Counter
    print()
    print(f"章证总数: {len(items)}")
    print(f"  章: {len([i for i in items if i['资料类别']=='章'])}")
    print(f"  执照: {len([i for i in items if i['资料类别']=='执照'])}")
    print(f"  在外部: {len([i for i in items if i['状态']=='在外部'])}")
    print(f"位置节点: {len([l for l in locations if l['类型']=='位置'])}")
    print(f"人员节点: {len([l for l in locations if l['类型']=='人员'])}")
    print(f"初始轨迹: {len(movements)}")
    return 0

if __name__ == "__main__":
    sys.exit(main())